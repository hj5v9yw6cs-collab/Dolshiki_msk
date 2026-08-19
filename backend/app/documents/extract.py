"""Извлечение текста из загруженных файлов.

Нужно для классификации и вытаскивания реквизитов. Сканы без текстового
слоя здесь не распознаются: OCR — отдельная тяжёлая история, и пока такие
файлы просто уходят на ручную проверку с честной пометкой.
"""

from __future__ import annotations

import logging
import re
import zipfile
from pathlib import Path

logger = logging.getLogger(__name__)

# Больше первых страниц читать незачем: тип документа и реквизиты
# определяются по началу, а разбор толстых файлов тормозит загрузку.
MAX_PDF_PAGES = 6
MAX_CHARS = 40_000


def _from_pdf(path: Path) -> str:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    chunks = []
    for page in reader.pages[:MAX_PDF_PAGES]:
        try:
            chunks.append(page.extract_text() or "")
        except Exception as exc:  # повреждённая страница не должна ронять загрузку
            logger.warning("Не удалось прочитать страницу PDF %s: %s", path.name, exc)
    return "\n".join(chunks)


def _from_docx(path: Path) -> str:
    from docx import Document as DocxDocument

    document = DocxDocument(str(path))
    parts = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.extend(cell.text for cell in row.cells)
    return "\n".join(parts)


def _from_doc(path: Path) -> str:
    """Текст из старого .doc (Word 97 — тот, что делает и WPS Office).

    Полноценного разбора формата тут нет: он громоздкий, а нам нужен только
    текст для определения типа и реквизитов. Поток WordDocument хранит
    символы либо как UTF-16, либо как однобайтовую кириллицу, поэтому
    пробуем оба варианта и берём тот, где кириллицы больше.
    """
    import olefile

    with olefile.OleFileIO(str(path)) as ole:
        if not ole.exists("WordDocument"):
            return ""
        raw = ole.openstream("WordDocument").read()

    best = ""
    for encoding in ("utf-16-le", "cp1251"):
        decoded = raw.decode(encoding, errors="ignore")
        # Форматирование внутри потока даёт мусор между словами: оставляем
        # только осмысленные последовательности букв, цифр и знаков.
        parts = re.findall(r"[А-Яа-яЁёA-Za-z0-9][^\x00-\x08\x0b\x0c\x0e-\x1f]{6,}", decoded)
        candidate = "\n".join(part.strip() for part in parts)
        if _cyrillic(candidate) > _cyrillic(best):
            best = candidate
    return best


def _cyrillic(text: str) -> int:
    return sum(1 for char in text if "а" <= char.lower() <= "я")


def _from_text(path: Path) -> str:
    for encoding in ("utf-8", "cp1251"):
        try:
            return path.read_text(encoding=encoding)
        except (UnicodeDecodeError, ValueError):
            continue
    return ""


def extract_text(path: Path, filename: str = "") -> str:
    """Текст документа или пустая строка, если извлечь нечего.

    Ошибки разбора не пробрасываются: файл всё равно должен сохраниться,
    просто без автоматического определения типа.
    """
    suffix = Path(filename or path.name).suffix.lower()
    try:
        if suffix == ".pdf":
            text = _from_pdf(path)
        elif suffix == ".doc":
            text = _from_doc(path)
        elif suffix in (".docx", ".dotx"):
            text = _from_docx(path)
        elif suffix in (".txt", ".rtf", ".csv"):
            text = _from_text(path)
        else:
            return ""
    except (zipfile.BadZipFile, OSError, ValueError) as exc:
        logger.warning("Не удалось извлечь текст из %s: %s", filename or path.name, exc)
        return ""
    except Exception as exc:  # неизвестный формат внутри знакомого расширения
        logger.warning("Ошибка разбора %s: %s", filename or path.name, exc)
        return ""

    return text[:MAX_CHARS]


def looks_like_scan(path: Path, filename: str, text: str) -> bool:
    """PDF без текстового слоя — почти наверняка скан."""
    suffix = Path(filename or path.name).suffix.lower()
    return suffix == ".pdf" and len(text.strip()) < 40
