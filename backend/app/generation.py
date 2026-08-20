"""Сборка претензии и искового заявления по шаблонам.

Тексты шаблонов лежат в `content/templates.json` и правятся юристом без
программиста — как тексты сайта и юридический конфиг. Здесь только
подстановка значений из карточки дела и сборка файла .docx.

Незаполненное поле не превращается в пустоту: в тексте остаётся видимая
пометка, а список пропусков возвращается вызывающему коду, чтобы юрист
узнал о них до отправки, а не из ответа суда.
"""

from __future__ import annotations

import io
import json
import re
import threading
from decimal import Decimal
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

from docx import Document as DocxDocument
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt

from .core.numbers import money_to_words
from .presenters import format_money

TEMPLATES_PATH = Path(__file__).resolve().parents[1] / "content" / "templates.json"

PLACEHOLDER_RE = re.compile(r"\{([a-z_]+)\}")

# Чем заменяется поле, которого в деле нет. Подчёркивания видно и на бумаге,
# и на экране: пустое место в готовом документе заметить труднее.
BLANK = "___________"

MONTHS_GENITIVE = (
    "января", "февраля", "марта", "апреля", "мая", "июня",
    "июля", "августа", "сентября", "октября", "ноября", "декабря",
)


class TemplateError(Exception):
    """Шаблоны не читаются или не проходят проверку."""


class TemplateStore:
    """Шаблоны с перечитыванием по времени изменения файла.

    Тот же приём, что у юридического конфига: правка вступает в силу без
    перезапуска, а сломанный файл не роняет систему — продолжает работать
    предыдущая исправная версия.
    """

    def __init__(self, path: Path) -> None:
        self._path = path
        self._lock = threading.Lock()
        self._loaded: Optional[Dict[str, Any]] = None
        self._mtime: float = 0.0

    def get(self) -> Dict[str, Any]:
        with self._lock:
            try:
                mtime = self._path.stat().st_mtime
            except OSError as error:
                if self._loaded is not None:
                    return self._loaded
                raise TemplateError(f"Файл шаблонов не найден: {self._path}") from error

            if self._loaded is None or mtime != self._mtime:
                try:
                    raw = json.loads(self._path.read_text(encoding="utf-8"))
                except (OSError, json.JSONDecodeError) as error:
                    if self._loaded is not None:
                        return self._loaded
                    raise TemplateError(f"Шаблоны не читаются: {error}") from error
                self._loaded = {
                    code: value for code, value in raw.items() if not code.startswith("_")
                }
                self._mtime = mtime
            return self._loaded


store = TemplateStore(TEMPLATES_PATH)


def catalog() -> List[dict]:
    """Список шаблонов для интерфейса."""
    return [
        {"code": code, "title": template.get("title", code),
         "doc_type": template.get("doc_type", "other")}
        for code, template in store.get().items()
    ]


# ---------------------------------------------------------------------------
# Значения для подстановки
# ---------------------------------------------------------------------------


def date_words(value) -> str:
    """«24 мая 2025 года» — так даты пишут в документах, а не 24.05.2025."""
    if value is None:
        return ""
    return f"{value.day} {MONTHS_GENITIVE[value.month - 1]} {value.year} года"


def _money(value) -> str:
    """Сумма без знака рубля.

    В документе валюту называет текст рядом — «составляет 15 613 672,30
    (Пятнадцать миллионов…) рубля 30 копеек». Значок ₽ там лишний и
    выглядит как опечатка.
    """
    return format_money(value) if value is not None else ""


def calculation_text(result: Optional[dict]) -> str:
    """Расчёт неустойки словами — абзац, как в претензии.

    Берётся из того же расчёта, что уходит приложением к иску: одна и та же
    сумма в тексте и в приложении, потому что источник один.
    """
    if not result:
        return ""

    period = result.get("period") or {}
    lines = [
        f"За период с {period.get('start_display', '')} по {period.get('end_display', '')} "
        f"просрочка передачи объекта долевого строительства составляет "
        f"{period.get('days_display', '')}."
    ]

    for segment in result.get("segments", []):
        lines.append("Расчёт: " + segment.get("formula", ""))

    excluded = result.get("excluded") or []
    if excluded:
        parts = [
            f"{item.get('start_display', '')} — {item.get('end_display', '')} ({item.get('basis', '')})"
            for item in excluded
        ]
        lines.append("Из расчёта исключены периоды, на которые неустойка не начисляется: "
                     + "; ".join(parts) + ".")

    return " ".join(lines)


def duty_values(case, result: Optional[dict]) -> Dict[str, str]:
    """Строка про госпошлину в шапке иска.

    Дольщик — потребитель, и чаще всего пошлину не платит вовсе. Написать
    в шапке «Госпошлина: 0 руб.» нельзя: это выглядит как незаполненное
    поле, а не как льгота, и суд вправе оставить иск без движения.
    Поэтому в шапку идёт готовая фраза, а не число.
    """
    from .core.duty import duty_for_claim

    lines = (result or {}).get("lines") or []
    if lines:
        claim_value = sum(
            (Decimal(line["amount"]) for line in lines
             if line.get("code") in ("neustoyka", "repair_cost")),
            Decimal(0),
        )
        duty = duty_for_claim(claim_value)
        amount, exempt = duty.amount, duty.exempt
    elif case.duty is not None:
        # Расчёта нет, но пошлину вписали руками — доверяем ей.
        amount, exempt = case.duty, case.duty == 0
    else:
        return {"duty": "", "duty_line": ""}

    exempt_text = (
        "истец освобождён от уплаты (пункт 3 статьи 17 Закона Российской "
        "Федерации «О защите прав потребителей», подпункт 4 пункта 2 "
        "статьи 333.36 Налогового кодекса Российской Федерации)"
    )
    return {
        "duty": _money(amount) if amount else "",
        "duty_line": exempt_text if exempt else f"{_money(amount)} руб.",
    }


def firm_values() -> Dict[str, str]:
    """Реквизиты самой практики — исполнителя и представителя.

    Берутся из того же файла, что и контакты сайта: реквизиты не должны
    жить в двух местах и расходиться. Меняются там же, где меняются
    контакты, — в content/site.json.
    """
    try:
        from .site.content import store as site_store

        company = site_store.get().company
    except Exception:  # сайт не обязан быть настроен, чтобы собрать документ
        return {}

    return {
        "firm_name": company.get("name", ""),
        "firm_legal_name": company.get("legal_name", ""),
        "firm_inn": str(company.get("inn", "")),
        "firm_address": company.get("address", ""),
        "firm_phone": company.get("phone", ""),
        "firm_email": company.get("email", ""),
    }


def case_values(case, calculation: Optional[dict] = None, today=None) -> Dict[str, str]:
    """Все подстановки для шаблона из одной карточки дела."""
    from datetime import date as date_type

    today = today or date_type.today()
    client = case.client
    developer = case.developer
    result = (calculation or {})

    total = result.get("total")
    total_decimal = Decimal(total) if total is not None else None

    values: Dict[str, str] = {
        "case_number": case.number or "",
        "today": today.strftime("%d.%m.%Y"),
        "today_words": date_words(today),

        "client_name": (client.full_name if client else "") or "",
        "client_phone": (client.phone if client else "") or "",
        "client_birth_date": client.birth_date.strftime("%d.%m.%Y") if client and client.birth_date else "",
        "client_passport": (client.passport if client else "") or "",
        "client_snils": (client.snils if client else "") or "",
        "client_inn": (client.inn if client else "") or "",
        "client_address": (client.address if client else "") or "",

        "developer_name": (developer.name if developer else "") or "",
        "developer_inn": (developer.inn if developer else "") or "",
        "developer_ogrn": (developer.ogrn if developer else "") or "",
        "developer_address": (developer.address if developer else "") or "",

        "contract_number": case.contract_number or "",
        "contract_date": case.contract_date.strftime("%d.%m.%Y") if case.contract_date else "",
        "contract_date_words": date_words(case.contract_date),
        "contract_price": _money(case.contract_price),
        "contract_price_words": money_to_words(case.contract_price) if case.contract_price else "",

        "due_date": case.due_date.strftime("%d.%m.%Y") if case.due_date else "",
        "due_date_words": date_words(case.due_date),
        "claim_sent_words": date_words(case.claim_sent_on),

        "apartment": case.apartment or "",
        # Десятичная запятая, а не точка: документ русский.
        "area": str(case.area).replace(".", ",") if case.area is not None else "",
        "object_address": case.object_address or "",
        "project": case.project or "",

        "court_name": case.court_name or "",
        "calculation_text": calculation_text(result),
        "total": _money(total_decimal),
        "total_words": money_to_words(total_decimal) if total_decimal is not None else "",

        # Размер морального вреда зависит от обстоятельств дела, поэтому
        # это поле карточки, а не расчёт: юрист вписывает сумму сам.
        "moral_damage": _money(case.moral_damage),
    }
    values.update(duty_values(case, result))
    values.update(firm_values())
    return values


# ---------------------------------------------------------------------------
# Подстановка и сборка файла
# ---------------------------------------------------------------------------


def fill(text: str, values: Dict[str, str], missing: set) -> str:
    def replace(match: re.Match) -> str:
        name = match.group(1)
        value = values.get(name, "")
        if not value:
            missing.add(name)
            return BLANK
        return value

    return PLACEHOLDER_RE.sub(replace, text)


def _paragraph(document, text: str, *, align=None, bold=False, space_after=6):
    paragraph = document.add_paragraph()
    run = paragraph.add_run(text)
    run.bold = bold
    if align is not None:
        paragraph.alignment = align
    paragraph.paragraph_format.space_after = Pt(space_after)
    return paragraph


def render(template_code: str, values: Dict[str, str]) -> Tuple[bytes, List[str], str]:
    """Собирает .docx по шаблону. Возвращает файл, список пропусков и имя."""
    templates = store.get()
    template = templates.get(template_code)
    if template is None:
        raise TemplateError(f"Шаблон «{template_code}» не найден.")

    missing: set = set()
    document = DocxDocument()

    style = document.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)

    for section in document.sections:
        section.left_margin = Cm(3)
        section.right_margin = Cm(1.5)
        section.top_margin = Cm(2)
        section.bottom_margin = Cm(2)

    for line in template.get("header_right", []):
        _paragraph(document, fill(line, values, missing),
                   align=WD_ALIGN_PARAGRAPH.RIGHT, space_after=0)

    document.add_paragraph()
    for line in template.get("heading", []):
        _paragraph(document, fill(line, values, missing),
                   align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, space_after=2)
    document.add_paragraph()

    for line in template.get("body", []):
        text = fill(line, values, missing).strip()
        if text:
            _paragraph(document, text, align=WD_ALIGN_PARAGRAPH.JUSTIFY)

    if template.get("demands_title"):
        _paragraph(document, template["demands_title"],
                   align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, space_after=8)
    for index, line in enumerate(template.get("demands", []), start=1):
        _paragraph(document, f"{index}. " + fill(line, values, missing),
                   align=WD_ALIGN_PARAGRAPH.JUSTIFY)

    for line in template.get("footer", []):
        _paragraph(document, fill(line, values, missing), align=WD_ALIGN_PARAGRAPH.JUSTIFY)

    if template.get("attachments"):
        document.add_paragraph()
        _paragraph(document, template.get("attachments_title", "Приложения:"), bold=True)
        for index, line in enumerate(template["attachments"], start=1):
            _paragraph(document, f"{index}. " + fill(line, values, missing), space_after=2)

    if template.get("signature"):
        document.add_paragraph()
        _paragraph(document, fill(template["signature"], values, missing))

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue(), sorted(missing), template.get("filename", template_code)
