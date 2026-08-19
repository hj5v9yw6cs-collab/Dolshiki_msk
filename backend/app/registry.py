"""Выгрузка реестра дел в Excel.

База — источник истины, файл — снимок для отчётности и для тех, кому
удобнее смотреть в таблице. Обратной загрузки нет намеренно: правки в
Excel не вернутся в систему и разъедутся с ней.
"""

from __future__ import annotations

from datetime import date, datetime
from decimal import Decimal
from io import BytesIO
from typing import Iterable, List, Optional, Sequence

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .casework import REGIONS, SERVICE_TYPES, next_deadline, stage_title

MONEY_FORMAT = '# ##0.00\\ "₽"'
DATE_FORMAT = "DD.MM.YYYY"

HEADER_FILL = PatternFill("solid", fgColor="F2F2F5")
OVERDUE_FILL = PatternFill("solid", fgColor="FDE8E8")
SOON_FILL = PatternFill("solid", fgColor="FFF7E0")


def _money(value: Optional[Decimal]):
    return float(value) if value is not None else None


# Excel считает формулой всё, что начинается с этих знаков, и выполняет её
# при открытии файла. Имя и объект в заявке пишет посетитель сайта, то есть
# кто угодно, — и его строка попадает в реестр, который открывает юрист.
# Апостроф перед значением делает ячейку текстом и обезвреживает подстановку.
FORMULA_STARTERS = ("=", "+", "-", "@", "\t", "\r")


def _safe(value):
    if isinstance(value, str) and value.startswith(FORMULA_STARTERS):
        return "'" + value
    return value


def _row(sheet, values) -> None:
    sheet.append([_safe(value) for value in values])


def _sheet(workbook: Workbook, title: str, columns: Sequence[tuple[str, int]]):
    sheet = workbook.create_sheet(title)
    _row(sheet, [name for name, _ in columns])
    for index, (_, width) in enumerate(columns, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
        cell = sheet.cell(row=1, column=index)
        cell.font = Font(bold=True, size=10)
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(columns))}1"
    return sheet


def _cases_sheet(workbook: Workbook, cases: Sequence, today: date, with_money: bool) -> None:
    columns = [
        ("№ дела", 12), ("Заведено", 12), ("Клиент", 26), ("Телефон", 18),
        ("Регион", 20), ("ЖК / объект", 24), ("Застройщик", 22), ("Услуга", 26),
        ("Стадия", 18), ("Юрист", 20), ("Цена ДДУ", 16), ("Заявлено", 16),
        ("Присуждено", 16), ("Получено", 16),
    ]
    if with_money:
        columns.append(("Гонорар", 14))
    columns += [
        ("Суд", 26), ("№ дела в суде", 18), ("Заседание", 13),
        ("Ближайший срок", 22), ("Дата срока", 13), ("Осталось дней", 14),
        ("Документов", 12), ("Комплект", 12),
    ]

    sheet = _sheet(workbook, "Дела", columns)

    for case in cases:
        deadline = next_deadline(case, today)
        documents = list(case.documents)
        confirmed = {item.doc_type for item in documents if not item.needs_review}

        from .documents import build_checklist
        checklist = build_checklist(case.stage, confirmed)

        row: List = [
            case.number,
            case.created_at.date(),
            case.client.full_name if case.client else "",
            case.client.phone if case.client else "",
            REGIONS.get(case.region, case.region),
            case.project or "",
            case.developer.name if case.developer else "",
            SERVICE_TYPES.get(case.service_type, case.service_type),
            stage_title(case.stage),
            case.lawyer.name if case.lawyer else "",
            _money(case.contract_price),
            _money(case.amount_claimed),
            _money(case.amount_awarded),
            _money(case.amount_received),
        ]
        if with_money:
            row.append(_money(case.fee))
        row += [
            case.court_name or "",
            case.court_case_number or "",
            case.next_hearing_on,
            deadline.title if deadline else "",
            deadline.due_on if deadline else None,
            deadline.days_left if deadline else None,
            len(documents),
            "готов" if checklist["ready"] else f"нет {checklist['missing_required']}",
        ]
        _row(sheet, row)

        current = sheet.max_row
        if deadline and deadline.is_overdue:
            fill = OVERDUE_FILL
        elif deadline and deadline.is_soon:
            fill = SOON_FILL
        else:
            fill = None
        if fill:
            for column in range(1, len(columns) + 1):
                sheet.cell(row=current, column=column).fill = fill

        for name, column_index in _columns_of(columns, ("Цена ДДУ", "Заявлено", "Присуждено", "Получено", "Гонорар")):
            sheet.cell(row=current, column=column_index).number_format = MONEY_FORMAT
        for name, column_index in _columns_of(columns, ("Заведено", "Заседание", "Дата срока")):
            sheet.cell(row=current, column=column_index).number_format = DATE_FORMAT


def _columns_of(columns: Sequence[tuple[str, int]], names: Iterable[str]):
    wanted = set(names)
    for index, (name, _) in enumerate(columns, start=1):
        if name in wanted:
            yield name, index


def _documents_sheet(workbook: Workbook, cases: Sequence) -> None:
    columns = [
        ("№ дела", 12), ("Клиент", 24), ("Тип документа", 28), ("Папка", 18),
        ("Имя файла", 44), ("Дата документа", 15), ("Загружен", 17),
        ("Кто загрузил", 20), ("Размер, КБ", 12), ("Проверен", 11),
    ]
    sheet = _sheet(workbook, "Документы", columns)

    for case in cases:
        for document in case.documents:
            _row(sheet, [
                case.number,
                case.client.full_name if case.client else "",
                _doc_title(document.doc_type),
                document.folder,
                document.stored_name,
                document.doc_date,
                document.created_at.replace(tzinfo=None),
                document.uploaded_by_name,
                max(document.size_bytes // 1024, 1),
                "нет" if document.needs_review else "да",
            ])
            current = sheet.max_row
            sheet.cell(row=current, column=6).number_format = DATE_FORMAT
            sheet.cell(row=current, column=7).number_format = "DD.MM.YYYY HH:MM"
            if document.needs_review:
                for column in range(1, len(columns) + 1):
                    sheet.cell(row=current, column=column).fill = SOON_FILL


def _doc_title(code: str) -> str:
    from .documents import type_title
    return type_title(code)


def _hearings_sheet(workbook: Workbook, cases: Sequence) -> None:
    columns = [
        ("Дата", 13), ("Время", 10), ("№ дела", 12), ("Клиент", 24),
        ("Суд", 30), ("№ дела в суде", 18), ("Юрист", 20),
    ]
    sheet = _sheet(workbook, "Заседания", columns)

    upcoming = sorted(
        (case for case in cases if case.next_hearing_on),
        key=lambda case: case.next_hearing_on,
    )
    for case in upcoming:
        _row(sheet, [
            case.next_hearing_on, "",
            case.number,
            case.client.full_name if case.client else "",
            case.court_name or "",
            case.court_case_number or "",
            case.lawyer.name if case.lawyer else "",
        ])
        sheet.cell(row=sheet.max_row, column=1).number_format = DATE_FORMAT


def build_registry(cases: Sequence, *, with_money: bool, today: Optional[date] = None) -> bytes:
    """Книга Excel с листами «Дела», «Документы» и «Заседания»."""
    today = today or date.today()
    workbook = Workbook()
    workbook.remove(workbook.active)

    _cases_sheet(workbook, cases, today, with_money)
    _documents_sheet(workbook, cases)
    _hearings_sheet(workbook, cases)

    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def registry_filename(today: Optional[date] = None) -> str:
    today = today or date.today()
    return f"reestr-del-{today.isoformat()}.xlsx"
